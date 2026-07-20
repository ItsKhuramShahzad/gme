#!/usr/bin/env python3
"""
make_abstract_docx.py — build the conference abstract as a clean MS-Word file.

Addresses every reviewer comment (G. Goldoni):
  * L -> N for the number of sites (n stays the cluster size, n <= N);
  * all equations written with Word super/subscript runs (no broken LaTeX);
  * C_GME defined explicitly (pure-state measure + convex-roof extension);
  * "weakest entanglement with respect to what?" -> stated: min over bipartitions
    of the bipartite concurrence built from the reduced-state purity;
  * no stray "L" anywhere.

Run:  python docs/make_abstract_docx.py
Out:  docs/abstract_GME_Hubbard.docx
"""
import os
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

BODY_FONT = "Times New Roman"
MATH_FONT = "Cambria Math"
BODY_SIZE = 11

doc = Document()
for s in doc.sections:
    s.left_margin = s.right_margin = Pt(54)     # ~0.75"
    s.top_margin = s.bottom_margin = Pt(54)

st = doc.styles["Normal"]
st.font.name = BODY_FONT
st.font.size = Pt(BODY_SIZE)
st.paragraph_format.space_after = Pt(8)
st.paragraph_format.line_spacing = 1.15


def _run(p, text, *, italic=False, bold=False, sup=False, sub=False, math=False):
    r = p.add_run(text)
    r.font.name = MATH_FONT if math else BODY_FONT
    r.font.size = Pt(BODY_SIZE)
    r.italic = italic
    r.bold = bold
    r.font.superscript = sup
    r.font.subscript = sub
    return r


def para(*tokens, align=None, space_after=None):
    """tokens: str (plain) or dict(text=..., italic/bold/sup/sub/math=...)."""
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    if space_after is not None:
        p.paragraph_format.space_after = Pt(space_after)
    for tk in tokens:
        if isinstance(tk, str):
            _run(p, tk)
        else:
            _run(p, tk.pop("text"), **tk)
    return p


def eq(*tokens):
    """Centered display equation (math font)."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    for tk in tokens:
        d = dict(tk); d["math"] = True
        _run(p, d.pop("text"), **d)
    return p


def I(t):   return {"text": t, "italic": True}
def P(t):   return {"text": t}
def SUP(t): return {"text": t, "sup": True}
def SUB(t): return {"text": t, "sub": True}
def IB(t):  return {"text": t, "italic": True, "bold": True}


# ── Title ────────────────────────────────────────────────────────────────────
h = doc.add_paragraph()
h.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = h.add_run("Estimation of Genuine Multipartite Entanglement "
              "in Half-Filled Hubbard Rings")
r.bold = True
r.font.name = BODY_FONT
r.font.size = Pt(14)

a = doc.add_paragraph()
a.alignment = WD_ALIGN_PARAGRAPH.CENTER
ra = a.add_run("Khuram Shahzad and Guido Goldoni")
ra.font.name = BODY_FONT
ra.font.size = Pt(11)
aff = doc.add_paragraph()
aff.alignment = WD_ALIGN_PARAGRAPH.CENTER
raf = aff.add_run("University of Modena and Reggio Emilia (UNIMORE), "
                  "Department of Physics, Informatics and Mathematics")
raf.font.name = BODY_FONT
raf.font.size = Pt(10)
raf.italic = True
aff.paragraph_format.space_after = Pt(14)

# ── P1 — motivation ──────────────────────────────────────────────────────────
para(
    "Genuine multipartite entanglement (GME) is a central resource for quantum "
    "technologies and a sensitive probe of many-body correlations. In "
    "itinerant-electron systems it is hard to quantify, because charge "
    "fluctuations, subsystem reduction, and temperature produce mixed spin "
    "states. Here we investigate spin GME in finite, half-filled Hubbard "
    "rings.")

# ── P2 — system + Hamiltonian ────────────────────────────────────────────────
para("We solve the Hubbard Hamiltonian,")
eq(I("H"), P(" = −t "), P("∑"), SUB("⟨ij⟩,σ"),
   P(" ("), I("c"), SUP("†"), SUB("iσ"), I("c"), SUB("jσ"),
   P(" + H.c.) + U "), P("∑"), SUB("i"), P(" "),
   I("n"), SUB("i↑"), P(" "), I("n"), SUB("i↓"), P(" ,"))
para(
    {"text": "by exact diagonalization for rings of "},
    I("N"), P(" = 2–8 sites at half filling ("), I("N"),
    P(" electrons, periodic boundary conditions), in both the "),
    I("S"), SUB("z"), P(" = 0 sector and the full Hilbert space. Here "),
    I("t"), P(" is the nearest-neighbour hopping and "), I("U"),
    P(" the on-site repulsion."))

# ── P3 — the reduced density matrix (define exactly what is calculated) ───────
para(
    {"text": "From a correlated eigenstate "}, I("Ψ"),
    P(" we choose a cluster of "), I("n"), P(" sites ("), I("n"),
    P(" ≤ "), I("N"),
    P(") and form their reduced density matrix by tracing out the remaining "),
    I("N"), P(" − "), I("n"), P(" sites,"))
eq(I("ρ"), SUB("n"), P(" = Tr"), SUB("N−n"),
   P(" |"), I("Ψ"), P("⟩⟨"), I("Ψ"), P("| ."))
para(
    "Projecting each selected site onto single occupation (projector ",
    I("P"), P(") and renormalizing yields the "), I("n"),
    P("-spin reduced density matrix"))
eq(I("Γ"), SUP("(n)"), P(" = "), I("P"), P(" "), I("ρ"), SUB("n"),
   P(" "), I("P"), P("  /  Tr("), I("P"), P(" "), I("ρ"), SUB("n"),
   P(" "), I("P"), P(") ,"))
para(
    {"text": "a 2"}, SUP("n"), P(" × 2"), SUP("n"),
    P(" matrix on the local spin configurations {↑,↓}"), SUP("n"),
    P(", with populations on the diagonal and spin coherences off it. This "
      "isolates the spin degrees of freedom from the charge while preserving "
      "the fermionic symmetries; "),
    I("Γ"), SUP("(n)"),
    P(" is Hermitian, positive semidefinite, and unit-trace. "
      "Finite-temperature states follow as Gibbs mixtures of the eigenstate "
      "matrices."))

# ── P4 — the GME measure (define C_GME; "with respect to what") ───────────────
para(
    "For a pure state, GME is quantified by the genuinely multipartite "
    "concurrence — the weakest bipartite entanglement over all bipartitions of "
    "the cluster into two parts ",
    I("A"), P(", "), I("B"), P(":"))
eq(I("C"), SUB("GME"), P("("), I("φ"), P(") = min"), SUB("A|B"),
   P(" √[ 2 ( 1 − Tr "), I("ρ"), SUB("A"), SUP("2"),
   P(" ) ] ,"))
para(
    {"text": "where "}, I("ρ"), SUB("A"),
    P(" is the reduced state of part "), I("A"),
    P(", minimized over all bipartitions of the "), I("n"),
    P(" spins. The extracted and thermal states are generally mixed, so their "
      "GME is defined through the convex-roof extension,"))
eq(I("C"), SUB("GME"), P("("), I("ρ"), P(") = min "), P("∑"),
   SUB("i"), P(" "), I("p"), SUB("i"), P(" "), I("C"), SUB("GME"),
   P("("), I("φ"), SUB("i"), P(") ,"))
para(
    {"text": "the minimum mean concurrence over all decompositions "},
    I("ρ"), P(" = ∑"), SUB("i"), P(" "), I("p"), SUB("i"),
    P(" |"), I("φ"), SUB("i"), P("⟩⟨"), I("φ"), SUB("i"),
    P("|. It is optimized by a generalized conjugate-gradient method on the "
      "unitary group, cross-checked by a Stiefel-manifold parametrization."))

# ── P5 — results ─────────────────────────────────────────────────────────────
para(
    {"text": "We map "}, I("C"), SUB("GME"),
    P(" as a function of the cluster size "), I("n"),
    P(", the coupling ratio "), I("t"), P("/"), I("U"),
    P(", the temperature "), I("T"), P(", and the ring size "), I("N"),
    P(". Ground-state GME generally grows with cluster size, approaching its "
      "maximum for the whole ring. Temperature suppresses it on a scale set "
      "by the spectral gap, so entanglement survives to higher "),
    I("T"), P(" for larger "), I("t"), P("/"), I("U"),
    P("; near separability, thermal population of entangled excited states "
      "can instead activate GME. For "),
    I("n"), P(" = 2 the measure reduces exactly to the Wootters concurrence, "
      "an analytic benchmark."))

# ── P6 — closing ─────────────────────────────────────────────────────────────
para(
    "This framework links correlated electronic structure to effective spin "
    "entanglement, clarifying how genuine multipartite correlations respond to "
    "interactions, subsystem reduction, and temperature; the same construction "
    "extends to molecular π systems such as benzene.")

out = os.path.join(os.path.dirname(__file__), "abstract_GME_Hubbard.docx")
doc.save(out)

# crude prose word count (excludes centered equations)
words = 0
for p in doc.paragraphs:
    if p.alignment == WD_ALIGN_PARAGRAPH.CENTER:
        continue
    words += len(p.text.split())
print(f"saved: {out}")
print(f"approx prose word count (equations excluded): {words}")
